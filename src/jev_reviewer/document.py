from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from pypdf import PdfReader

from .models import Document


SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(path: Path) -> tuple[str, dict[str, Any]]:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text() or ""
        pages.append(f"\n[PAGE {i}]\n{extracted}")
    meta = dict(reader.metadata or {})
    return _clean_text("\n".join(pages)), {"pages": len(reader.pages), "pdf_metadata": meta}


def _docx(path: Path) -> tuple[str, dict[str, Any]]:
    doc = WordDocument(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    props = doc.core_properties
    meta = {
        "author": props.author,
        "last_modified_by": props.last_modified_by,
        "created": props.created.isoformat() if props.created else None,
        "modified": props.modified.isoformat() if props.modified else None,
        "revision": props.revision,
        "title": props.title,
    }
    return _clean_text("\n".join(parts)), {"docx_metadata": meta}


def load_document(path: Path) -> Document:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix}. Supported: {', '.join(sorted(SUPPORTED))}")
    if suffix == ".pdf":
        text, metadata = _pdf(path)
        kind = "pdf"
    elif suffix == ".docx":
        text, metadata = _docx(path)
        kind = "docx"
    else:
        text = _clean_text(path.read_text(encoding="utf-8", errors="replace"))
        metadata = {}
        kind = suffix.lstrip(".")
    if not text:
        raise ValueError("I could not extract text from this file. Scanned PDFs are not OCRed in V1.")
    return Document(name=path.name, path=path, text=text, kind=kind, metadata=metadata)


def ooxml_generator_hints(path: Path) -> list[str]:
    if path.suffix.lower() != ".docx":
        return []
    hints: list[str] = []
    try:
        with zipfile.ZipFile(path) as zf:
            core = zf.read("docProps/core.xml").decode("utf-8", errors="ignore") if "docProps/core.xml" in zf.namelist() else ""
            app = zf.read("docProps/app.xml").decode("utf-8", errors="ignore") if "docProps/app.xml" in zf.namelist() else ""
            joined = (core + "\n" + app).lower()
            for marker in ["python-docx", "libreoffice", "pandoc", "microsoft office", "microsoft word"]:
                if marker in joined:
                    hints.append(marker)
    except zipfile.BadZipFile:
        hints.append("invalid-ooxml-zip")
    return hints
